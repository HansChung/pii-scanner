"""抽不到文字的檔案（如掃描影像 PDF）要標為未完成查驗，不能顯示成無風險。"""
from __future__ import annotations

import io
import uuid
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image

from app.db import get_db
from app.report_export import build_excel, build_pdf, collect_report
from app.scanner import run_scan


def _image_only_pdf(path: Path) -> None:
    Image.new("RGB", (400, 200), "white").save(path, "PDF")


def _prepare_job(app, files: list[tuple[str, str, bytes | None]]) -> tuple[str, list[dict]]:
    """建立任務與檔案紀錄；files 為 (檔名, 副檔名, 內容)，內容為 None 時寫入影像 PDF。"""
    job_id = str(uuid.uuid4())
    upload_dir = Path(app.config["TEMP_UPLOAD_DIR"]) / job_id
    upload_dir.mkdir(parents=True)
    now = datetime.now(timezone.utc).isoformat()
    records = []
    db = get_db()
    db.execute(
        """
        INSERT INTO scan_jobs
        (id, status, progress, message, file_count, total_size, created_at, updated_at)
        VALUES (?, 'queued', 0, 'queued', ?, 0, ?, ?)
        """,
        (job_id, len(files), now, now),
    )
    for name, extension, content in files:
        file_id = str(uuid.uuid4())
        path = upload_dir / f"{file_id}-{name}"
        if content is None:
            _image_only_pdf(path)
        else:
            path.write_bytes(content)
        db.execute(
            """
            INSERT INTO files (id, job_id, original_name, extension, size, sha256, status, created_at)
            VALUES (?, ?, ?, ?, ?, 'hash', 'queued', ?)
            """,
            (file_id, job_id, name, extension, path.stat().st_size, now),
        )
        records.append({"file_id": file_id, "path": str(path), "extension": extension})
    db.commit()
    return job_id, records


def test_image_only_pdf_is_unreadable_not_no_risk(client, monkeypatch):
    monkeypatch.setattr("app.scanner.document_intelligence_enabled", lambda: False)
    app = client.application
    with app.app_context():
        job_id, records = _prepare_job(app, [("掃描名冊.pdf", ".pdf", None)])
        run_scan(job_id, records)
        db = get_db()
        job = db.execute("SELECT * FROM scan_jobs WHERE id = ?", (job_id,)).fetchone()
        file = db.execute("SELECT status, error FROM files WHERE job_id = ?", (job_id,)).fetchone()

    assert job["status"] == "completed"
    assert job["risk_level"] == "Unknown"
    assert "未完成查驗" in job["message"]
    assert file["status"] == "unreadable"
    assert "OCR" in file["error"]

    api_job = client.get(f"/api/jobs/{job_id}").get_json()
    assert api_job["files"][0]["status"] == "unreadable"


def test_unreadable_file_listed_in_reports(client, monkeypatch):
    monkeypatch.setattr("app.scanner.document_intelligence_enabled", lambda: False)
    app = client.application
    with app.app_context():
        job_id, records = _prepare_job(app, [("掃描名冊.pdf", ".pdf", None)])
        run_scan(job_id, records)
        report = collect_report(job_id)
        workbook = load_workbook(io.BytesIO(build_excel(report)))
        assert build_pdf(report).startswith(b"%PDF")

    summary = {row[0]: row[1] for row in workbook["查驗摘要"].iter_rows(values_only=True)}
    assert summary["整體風險"].startswith("未確認")
    assert "掃描名冊.pdf" in summary["未完成查驗檔案"]


def test_unreadable_file_does_not_lower_high_risk(client, monkeypatch):
    monkeypatch.setattr("app.scanner.document_intelligence_enabled", lambda: False)
    app = client.application
    with app.app_context():
        job_id, records = _prepare_job(
            app,
            [
                ("掃描名冊.pdf", ".pdf", None),
                ("名單.txt", ".txt", "身分證 A123456789".encode()),
            ],
        )
        run_scan(job_id, records)
        db = get_db()
        job = db.execute("SELECT risk_level, message FROM scan_jobs WHERE id = ?", (job_id,)).fetchone()
        statuses = {
            row["original_name"]: row["status"]
            for row in db.execute("SELECT original_name, status FROM files WHERE job_id = ?", (job_id,))
        }

    assert job["risk_level"] == "High"
    assert "1 個檔案" in job["message"]
    assert statuses == {"掃描名冊.pdf": "unreadable", "名單.txt": "completed"}


def test_ocr_that_returns_no_text_is_still_unreadable(client, monkeypatch):
    monkeypatch.setattr("app.scanner.document_intelligence_enabled", lambda: True)
    monkeypatch.setattr(
        "app.scanner.extract_with_document_intelligence", lambda path, heartbeat=None: []
    )
    app = client.application
    with app.app_context():
        job_id, records = _prepare_job(app, [("空白掃描.pdf", ".pdf", None)])
        run_scan(job_id, records)
        file = get_db().execute("SELECT status, error FROM files WHERE job_id = ?", (job_id,)).fetchone()

    assert file["status"] == "unreadable"
    assert file["error"].startswith("OCR 後仍未讀到")


def test_readable_file_without_findings_is_no_risk(client):
    app = client.application
    with app.app_context():
        job_id, records = _prepare_job(app, [("公告.txt", ".txt", "本週五停課一天。".encode())])
        run_scan(job_id, records)
        job = get_db().execute("SELECT risk_level, message FROM scan_jobs WHERE id = ?", (job_id,)).fetchone()

    assert job["risk_level"] == "None"
    assert job["message"] == "查驗完成"
