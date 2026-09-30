"""表格欄位標題感知：表頭在欄位、個資在儲存格時，也要抓得到姓名、學號等需要關鍵字的個資。"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

from app.extractors import extract_text
from app.pii_rules import detect_with_rules
from pii_scanner.scanners.spreadsheet_text import labeled_rows


def _findings(path: Path, extension: str):
    return [
        finding
        for chunk in extract_text(path, extension)
        for finding in detect_with_rules(chunk.text, chunk.location)
    ]


def _categories_at(findings, location: str) -> set[str]:
    return {finding.category for finding in findings if finding.location == location}


def test_xlsx_header_below_title_labels_each_cell(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "核定名單"
    ws.append(["114學年度獎助學金核定名單"])
    ws.append(["學號", "姓名", "身分證字號", "手機", "匯款帳號", "核定金額"])
    ws.append([411730000, "王小明", "A123456789", "0912-345-678", "700-1234-5678901", 20000])
    path = tmp_path / "名單.xlsx"
    wb.save(path)

    findings = _findings(path, ".xlsx")
    categories = _categories_at(findings, "工作表 核定名單 第 3 列")
    assert {"StudentNumber", "PersonName", "TaiwanNationalId", "MobilePhone", "BankAccount"} <= categories
    # 相鄰儲存格（身分證、手機）不應被串成信用卡號
    assert "CreditCard" not in categories


def test_docx_table_with_header(tmp_path):
    doc = Document()
    table = doc.add_table(rows=2, cols=3)
    for cell, text in zip(table.rows[0].cells, ["姓名", "出生日期", "緊急聯絡電話"]):
        cell.text = text
    for cell, text in zip(table.rows[1].cells, ["吳宜蓁", "2009/03/14", "0921-663-018"]):
        cell.text = text
    path = tmp_path / "報名表.docx"
    doc.save(path)

    categories = _categories_at(_findings(path, ".docx"), "表格 1 列 2")
    assert {"PersonName", "BirthDate", "MobilePhone"} <= categories


def test_docx_form_style_table_still_detected(tmp_path):
    # 表單式「欄名｜值｜欄名｜值」不能被誤當成表頭，否則下一列的生日會對不上欄名
    doc = Document()
    table = doc.add_table(rows=2, cols=4)
    for cell, text in zip(table.rows[0].cells, ["姓名", "王小明", "性別", "男"]):
        cell.text = text
    for cell, text in zip(table.rows[1].cells, ["出生日期", "2009/03/14", "電話", "0912-345-678"]):
        cell.text = text
    path = tmp_path / "個人資料表.docx"
    doc.save(path)

    findings = _findings(path, ".docx")
    names = [f.masked_text for f in findings if f.category == "PersonName"]
    assert names == ["王**"]
    assert "BirthDate" in _categories_at(findings, "表格 1 列 2")


def test_csv_with_header(tmp_path):
    path = tmp_path / "簽到表.csv"
    path.write_text(
        "學號,姓名,畢業系所,手機\n411730001,劉雅婷,教育科技學系,0918-775-230\n", encoding="utf-8-sig"
    )

    findings = _findings(path, ".csv")
    categories = _categories_at(findings, "CSV 第 2 列")
    assert {"StudentNumber", "PersonName", "MobilePhone"} <= categories
    # 表頭列「姓名、畢業系所」不應把「畢業系所」當成姓名
    assert not _categories_at(findings, "CSV 第 1 列")


def test_pptx_table_is_scanned_with_header(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    shape = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(6), Inches(1.5))
    for cell, text in zip(shape.table.rows[0].cells, ["學號", "姓名"]):
        cell.text = text
    for cell, text in zip(shape.table.rows[1].cells, ["411730002", "許文彬"]):
        cell.text = text
    path = tmp_path / "簡報.pptx"
    prs.save(path)

    categories = _categories_at(_findings(path, ".pptx"), "投影片 1 表格 1 列 2")
    assert {"StudentNumber", "PersonName"} <= categories


def test_labeled_rows_keep_original_row_numbers():
    rows = labeled_rows([["姓名", "手機"], [], ["王小明", "0912345678"]])
    assert [(row.row_number, row.is_header) for row in rows] == [(1, True), (3, False)]
    assert rows[1].text == "姓名：王小明\t手機：0912345678"


def test_value_rows_are_never_headers():
    rows = labeled_rows([["出生日期", "2009/03/14", "電話", "0912-345-678"], ["姓名", "王小明"]])
    assert not any(row.is_header for row in rows)
